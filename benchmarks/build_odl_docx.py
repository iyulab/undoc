"""Build the DOCX edition of opendataloader-bench and convert it with undoc.

There is no public DOCX-to-Markdown ground-truth dataset, so the published DOCX scores
are measured on a corpus built here. The content comes from the real-document ground
truth of opendataloader-bench (200 documents with headings, paragraphs and tables that
use merged cells). Each ground-truth file is typeset as a Word document using Word
semantics (Heading styles, real tables with merged cells), and the ground truth for the
DOCX edition is the same structure written back as Markdown. Scoring is done by that
benchmark's own evaluator, not by a metric of ours.

What it measures: whether declared document structure (headings, paragraphs, tables
with merged cells) survives conversion. It does not measure layout inference, because
the typeset documents carry their structure explicitly.

Typesetting rules:
  * `#` .. `######` lines become Heading 1 .. Heading 6.
  * A `<table>` block becomes a Word table (style "Table Grid"); rowspan and colspan
    become merged cells.
  * Every other run of lines separated by blank lines becomes one paragraph.
  * Image references are not content and are dropped.
  * The ground truth is the typeset structure written back as Markdown (paragraph
    line breaks are normalized to one line per paragraph).

Usage:
    python benchmarks/build_odl_docx.py <odl-bench-dir> <out-dir>
        writes <out-dir>/docx/*.docx and <out-dir>/ground-truth/markdown/*.md
    python benchmarks/build_odl_docx.py --convert <out-dir>
        writes <out-dir>/prediction/undoc/markdown/*.md using the installed undoc package

Dependencies: python-docx, beautifulsoup4, lxml (build); undoc (convert).
"""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

HEADING = re.compile(r"^(#{1,6})\s+(.*)$")
IMAGE = re.compile(r"^!\[[^\]]*\]\([^)]*\)\s*$")


def blocks(markdown: str):
    """Yield ('heading', level, text) | ('table', html) | ('para', text) in order."""
    lines = markdown.splitlines()
    i = 0
    para: list[str] = []

    def flush():
        if para:
            text = " ".join(s.strip() for s in para).strip()
            para.clear()
            if text:
                return ("para", text)
        return None

    while i < len(lines):
        line = lines[i]
        if line.strip().lower().startswith("<table"):
            if (p := flush()):
                yield p
            buf = [line]
            while "</table>" not in lines[i].lower() and i + 1 < len(lines):
                i += 1
                buf.append(lines[i])
            yield ("table", "\n".join(buf))
        elif (m := HEADING.match(line)):
            if (p := flush()):
                yield p
            yield ("heading", len(m.group(1)), m.group(2).strip())
        elif not line.strip() or IMAGE.match(line.strip()):
            if (p := flush()):
                yield p
        else:
            para.append(line)
        i += 1
    if (p := flush()):
        yield p


def table_grid(html: str):
    """Return (placed cells, row count, grid width); cells are (row, col, rowspan, colspan, text)."""
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(html, "lxml")
    rows = []
    for tr in soup.find_all("tr"):
        cells = []
        for td in tr.find_all(["td", "th"]):
            text = re.sub(r"\s+", " ", td.get_text(" ")).strip()
            cells.append((text, int(td.get("rowspan", 1) or 1), int(td.get("colspan", 1) or 1)))
        rows.append(cells)
    # Lay the cells onto a grid to learn its width (rowspans occupy later rows).
    occupied: set[tuple[int, int]] = set()
    width = 0
    placed = []
    for r, cells in enumerate(rows):
        c = 0
        for text, rs, cs in cells:
            while (r, c) in occupied:
                c += 1
            placed.append((r, c, rs, cs, text))
            for dr in range(rs):
                for dc in range(cs):
                    occupied.add((r + dr, c + dc))
            c += cs
            width = max(width, c)
    return placed, len(rows), width


def build(odl_dir: Path, out_dir: Path) -> int:
    """Typeset every ground-truth file of `odl_dir`; return the number of documents built."""
    import docx

    docx_dir = out_dir / "docx"
    gt_dir = out_dir / "ground-truth" / "markdown"
    docx_dir.mkdir(parents=True, exist_ok=True)
    gt_dir.mkdir(parents=True, exist_ok=True)
    for gt_path in sorted((odl_dir / "ground-truth" / "markdown").glob("*.md")):
        doc = docx.Document()
        out_md: list[str] = []
        for block in blocks(gt_path.read_text(encoding="utf-8")):
            if block[0] == "heading":
                _, level, text = block
                doc.add_heading(text, level=level)
                out_md.append(f"{'#' * level} {text}")
            elif block[0] == "para":
                doc.add_paragraph(block[1])
                out_md.append(block[1])
            else:
                placed, nrows, ncols = table_grid(block[1])
                if not nrows or not ncols:
                    continue
                table = doc.add_table(rows=nrows, cols=ncols)
                table.style = "Table Grid"
                for r, c, rs, cs, text in placed:
                    r2, c2 = min(r + rs - 1, nrows - 1), min(c + cs - 1, ncols - 1)
                    cell = table.cell(r, c)
                    if (r2, c2) != (r, c):
                        cell = cell.merge(table.cell(r2, c2))
                    cell.text = text
                out_md.append(block[1].strip())
        doc.save(docx_dir / f"{gt_path.stem}.docx")
        (gt_dir / gt_path.name).write_text("\n\n".join(out_md) + "\n", encoding="utf-8")
    return len(list(docx_dir.glob("*.docx")))


def convert(out_dir: Path) -> int:
    """Convert <out_dir>/docx/*.docx with the installed undoc package; return the failure count.

    A document that fails to convert gets an empty prediction so that it scores zero.
    """
    import undoc

    md_dir = out_dir / "prediction" / "undoc" / "markdown"
    md_dir.mkdir(parents=True, exist_ok=True)
    failures = 0
    for path in sorted((out_dir / "docx").glob("*.docx")):
        try:
            md = undoc.parse_file(path).to_markdown()
        except Exception as exc:
            print(f"failed: {path.name}: {type(exc).__name__}: {exc}", file=sys.stderr)
            failures += 1
            md = ""
        (md_dir / f"{path.stem}.md").write_text(md, encoding="utf-8")
    return failures


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--convert", metavar="OUT_DIR", type=Path,
                    help="convert <OUT_DIR>/docx with the installed undoc package")
    ap.add_argument("paths", nargs="*", type=Path, help="<odl-bench-dir> <out-dir> (build mode)")
    args = ap.parse_args()
    if args.convert:
        if args.paths:
            ap.error("--convert takes only its out-dir")
        failures = convert(args.convert)
        print(f"converted with {failures} failure(s) into {args.convert / 'prediction' / 'undoc' / 'markdown'}")
        return 0
    if len(args.paths) != 2:
        ap.error("build mode needs <odl-bench-dir> <out-dir>")
    odl_dir, out_dir = args.paths
    count = build(odl_dir, out_dir)
    print(f"built {count} documents in {out_dir}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
